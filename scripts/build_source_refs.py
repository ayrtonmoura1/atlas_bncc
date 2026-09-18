"""Build a workbook-backed map from spreadsheet locations to their skill code."""
import json
import re
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parents[1]
BOOK = ROOT / 'BNCC_Foco.xlsx'
CODE = re.compile(r'(?<![A-Z0-9])(?:EF|EM|EI)\d{2}[A-Z]{2,3}\d{2,3}(?![A-Z0-9])')
CELL = re.compile(r'^(.*)!([A-Z]+\d+)$')

def codes(value):
    return list(dict.fromkeys(CODE.findall(str(value or '').upper())))

def col_number(letters):
    value = 0
    for char in letters:
        value = value * 26 + ord(char) - 64
    return value

wb = openpyxl.load_workbook(BOOK, read_only=True, data_only=True)
cells, rows, row_codes = {}, {}, {}
for sheet in wb.sheetnames:
    ws = wb[sheet]
    for row_number, row in enumerate(ws.iter_rows(), 1):
        row_values = [cell.value for cell in row]
        found = []
        for value in row_values:
            found.extend(codes(value))
        found = list(dict.fromkeys(found))
        if not found:
            continue
        row_key = f'{sheet}#{row_number}'
        row_codes[row_key] = found
        for cell in row:
            direct = codes(cell.value)
            if not direct:
                continue
            key = f'{sheet}!{cell.coordinate}'
            # In detailed maps, relationship/comment columns contain referenced
            # skills while column C identifies the skill owning the row.
            c_value = row[2].value if len(row) > 2 else None
            owner = codes(c_value)
            if owner and cell.column != 3:
                cells[key] = owner[0]
            else:
                cells[key] = direct[0]
        # Every populated row can be referred to by its row number. Prefer the
        # code in column C, then the first code in the row.
        c_value = row[2].value if len(row) > 2 else None
        owner = codes(c_value) or found
        rows[row_key] = owner[0]
    # Fill blank relationship/comment cells in a row through the row owner.
    for row_number in range(1, ws.max_row + 1):
        row_key = f'{sheet}#{row_number}'
        if row_key not in rows:
            continue
        for col in range(1, ws.max_column + 1):
            key = f'{sheet}!{openpyxl.utils.get_column_letter(col)}{row_number}'
            cells.setdefault(key, rows[row_key])
wb.close()

graph = json.loads((ROOT / 'dados/BNCC_Grafo.json').read_text(encoding='utf-8'))
locations = set()
for node in graph['nodes']:
    for mention in node.get('mentions', []):
        locations.add(f"{mention['sheet']}!{mention['cell']}")
    for record in node.get('records', []):
        if record.get('sheet') and record.get('row'):
            locations.add(f"{record['sheet']}#{record['row']}")
for edge in graph['edges'] + graph.get('unconfirmed_edges', []):
    for evidence in edge.get('evidence', []):
        if evidence.get('location'):
            locations.add(evidence['location'])
for issue in graph.get('issues', []):
    locations.update(re.findall(r'[^;\n]+![A-Z]+\d+', issue.get('source', '')))
    locations.update(re.findall(r'[^;\n]+![A-Z]+\d+', issue.get('value', '')))

unresolved = sorted(location for location in locations if (location not in cells and location not in rows))
mapping = {'cells': cells, 'rows': rows, 'unresolved': unresolved, 'source': 'BNCC_Foco.xlsx'}
graph['metadata']['source_reference_policy'] = 'Referências de aba/célula ou aba/linha são exibidas pelo código da habilidade correspondente à linha na BNCC_Foco.xlsx.'
graph['metadata']['source_references'] = mapping
(ROOT / 'dados/source_references.json').write_text(json.dumps(mapping, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
(ROOT / 'dados/BNCC_Grafo.json').write_text(json.dumps(graph, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')

data_path = ROOT / 'data.js'
source = data_path.read_text(encoding='utf-8')
data = json.loads(source[source.index('{'):].strip().rstrip(';'))
data['metadata']['source_reference_policy'] = graph['metadata']['source_reference_policy']
data['metadata']['source_references'] = mapping
data_path.write_text('window.BNCC_DATA = '+json.dumps(data, ensure_ascii=False, separators=(',', ':'))+';\n', encoding='utf-8')

explorer_path = ROOT / 'explorador.html'
explorer = explorer_path.read_text(encoding='utf-8')
start = explorer.index('const DATA=') + len('const DATA=')
end = explorer.index(';\nconst $=', start)
explorer = explorer[:start] + json.dumps(graph, ensure_ascii=False, separators=(',', ':')) + explorer[end:]
explorer_path.write_text(explorer, encoding='utf-8')
print(json.dumps({'locations': len(locations), 'cell_mappings': len(cells), 'row_mappings': len(rows), 'unresolved': unresolved[:30], 'unresolved_count': len(unresolved)}, ensure_ascii=False))
